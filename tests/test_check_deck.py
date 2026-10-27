import importlib.util
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "preza_check_deck", ROOT / "skills/preza/scripts/check_deck.py"
)
check_deck = importlib.util.module_from_spec(spec)
spec.loader.exec_module(check_deck)

W, H = Emu(12192000), Emu(6858000)  # 16:9


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # пустой макет


def text(s, x, y, w, h, body, size=20):
    box = s.shapes.add_textbox(
        Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h))
    )
    box.text_frame.word_wrap = True
    box.text_frame.text = body
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return box


def rect(s, x, y, w, h):
    shp = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(int(W * x)),
        Emu(int(H * y)),
        Emu(int(W * w)),
        Emu(int(H * h)),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    return shp


def full(s, sparse_first=False):
    """Плотный слайд: заголовок и сетка 3×2 карточек с текстом; sparse_first — в первой карточке одна строка."""
    text(s, 0.05, 0.04, 0.9, 0.1, 'Двое держат половину рынка', 32)
    for i in range(3):
        for j in range(2):
            x, y = 0.05 + i * 0.31, 0.18 + j * 0.4
            rect(s, x, y, 0.29, 0.37)
            if sparse_first and i == j == 0:
                text(s, x + 0.01, y + 0.01, 0.27, 0.05, 'Одна строка', 18)
            else:
                text(s, x + 0.01, y + 0.01, 0.27, 0.35, 'Пункт с объяснением ' * 9, 18)


def run(prs, **kw):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "deck.pptx"
        prs.save(path)
        return check_deck.check(path, skip=(), **kw)


class VoidTests(unittest.TestCase):
    def test_half_empty_slide_is_found_and_full_is_not(self):
        prs = deck()
        s = slide(prs)
        text(s, 0.05, 0.05, 0.4, 0.85, "Только левая половина " * 15, 20)
        full(slide(prs))
        res = run(prs)
        self.assertTrue(any(f["kind"] == "void" for f in res[1]))
        self.assertFalse(any(f["kind"] == "void" for f in res.get(2, [])))

    def test_background_picture_is_not_content(self):
        prs = deck()
        s = slide(prs)
        with tempfile.TemporaryDirectory() as tmp:
            png = Path(tmp) / "bg.png"
            from PIL import Image

            Image.new("RGB", (160, 90), (90, 120, 160)).save(png)
            s.shapes.add_picture(str(png), 0, 0, W, H)
            text(s, 0.35, 0.45, 0.3, 0.1, "Заголовок", 28)
            res = run(prs)
        self.assertTrue(any(f["kind"] == "void" for f in res[1]))

    def test_card_with_one_line_is_void_inside_container(self):
        prs = deck()
        full(slide(prs), sparse_first=True)
        full(slide(prs))
        res = run(prs)
        self.assertTrue(any(f['kind'] == 'card' for f in res[1]))
        self.assertFalse(any(f['kind'] in ('card', 'void') for f in res.get(2, [])))

    def test_accepted_and_skipped_slides_do_not_fail(self):
        prs = deck()
        text(slide(prs), 0.05, 0.05, 0.4, 0.2, "Мало", 20)
        res = run(prs, accept=(1,))
        self.assertTrue(
            all(f["level"] == "accepted" for f in res[1] if f["kind"] == "void")
        )
        self.assertEqual(check_deck.exit_code(res), 0)

    def test_group_children_use_group_transform(self):
        prs = deck()
        s = slide(prs)
        grp = s.shapes.add_group_shape()
        text(grp, 0.0, 0.0, 0.2, 0.1, "Внутри группы", 18)
        grp.left, grp.top = Emu(int(W * 0.6)), Emu(int(H * 0.7))
        boxes = check_deck.content_boxes(s, W, H)
        x0, y0, _, _ = boxes[0]
        self.assertAlmostEqual(x0, 0.6, places=2)
        self.assertAlmostEqual(y0, 0.7, places=2)


class TextTests(unittest.TestCase):
    def kinds(self, body, y=0.4, strict=False, table=False):
        prs = deck()
        s = slide(prs)
        full(s)
        if table:
            tbl = s.shapes.add_table(
                1,
                1,
                Emu(int(W * 0.1)),
                Emu(int(H * 0.02)),
                Emu(int(W * 0.2)),
                Emu(int(H * 0.05)),
            ).table
            tbl.cell(0, 0).text = body
        else:
            text(s, 0.4, y, 0.2, 0.05, body, 14)
        return [
            f["kind"] for f in run(prs, strict=strict).get(1, []) if f["kind"] != "void"
        ]

    def test_clean_texts_pass(self):
        for body in (
            "информации о компании",
            "РОССИИ",
            "модели ИИ меняют рынок",
            "сгенерированный денежный поток",
            "24/7",
        ):
            self.assertEqual(self.kinds(body), [], body)
        self.assertEqual(self.kinds("1/3", y=0.02, table=True), [])

    def test_service_marks_are_found(self):
        self.assertEqual(self.kinds("03/12", y=0.01), ["counter"])
        self.assertEqual(self.kinds("Изображение сгенерировано нейросетью"), ["stop"])
        self.assertEqual(self.kinds("смоделировано"), [])
        self.assertEqual(self.kinds("смоделировано", strict=True), ["stop"])


class TitleTests(unittest.TestCase):
    def kinds(self, **kw):
        prs = deck()
        first, second = slide(prs), slide(prs)
        text(first, 0.05, 0.05, 0.9, 0.08, "МИКРОЭКОНОМИКА · ТЕМА 5", 14)
        text(first, 0.05, 0.3, 0.9, 0.2, "Рынок платформ: итоги года", 44)
        text(first, 0.05, 0.8, 0.9, 0.08, "Алиев Тимур, Кузнец Иван 402 группа", 16)
        text(second, 0.05, 0.85, 0.9, 0.08, "Спасибо за внимание! Алиева, Кузнец", 16)
        found = {no: sorted(f["kind"] for f in fs if f["kind"] != "void") for no, fs in run(prs, **kw).items()}
        return {no: k for no, k in found.items() if k}  # пустоты тестовых слайдов здесь не предмет

    def test_title_service_line_and_authors_outside_title(self):
        self.assertEqual(self.kinds(authors=["Алиев", "Кузнец"]), {1: ["title"], 2: ["authors", "authors"]})

    def test_allow_title_and_no_authors(self):
        self.assertEqual(self.kinds(allow_title=True), {})


if __name__ == "__main__":
    unittest.main()
