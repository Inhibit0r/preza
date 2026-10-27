#!/usr/bin/env python3
"""Пустоты и служебные пометки на слайдах PPTX до показа пользователю.

    python3 check_deck.py deck.pptx [--render 5-проверка/roundN] [--skip 1,13] [--accept 7] [--empty 0.30] [--strict] [--allow REGEX]
                          [--authors "Фамилия,Фамилия"] [--allow-title] [--json]
    python3 check_deck.py --selftest

Пустота — самый большой прямоугольник слайда без содержимого (доля площади ≥ --empty) или пустая часть
карточки-контейнера. Содержимое: картинки (кроме фона ≥90 % слайда), графики, таблицы, мелкие фигуры
и занятая текстом площадь (оценка по кеглю и переносу; рамки wrap="none" — как есть).
--skip — слайды без проверки пустот (по умолчанию первый и последний), --accept — намеренные пустоты
из раскадровки. Стоп-список (stoplist.py) — предупреждение, при --strict (учебная сдача) — ошибка.
Титул (слайд 1): служебная строка курса, темы, вида работы или ссылки на главу и источник — ошибка, кроме
--allow-title (задание требует курс на титуле). --authors: фамилии авторов на слайдах со второго — ошибка.
Код выхода 1 — непринятая пустота, счётчик «NN/NN», номер слайда, служебная строка титула, фамилия автора
вне титула или стоп-фраза при --strict.
"""

import argparse
import json
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import stoplist  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

# Калибровка 27.10.2026 на учебных колодах с жалобами на пустоты и принятых без них: принятые — 0 находок;
# колода, названная «очень бедной», — 5 пустых слайдов и счётчики NN/20; «скудная» — 3 слайда. Пропуск: пустота ~1/3
# карточки. EMPTY 0.35 → 0.30: добавились пустые слайды и выводы узкой колонкой (правая треть слайда пуста);
# пустоты внутри панелей по геометрии PPTX не видны — их ловит осмотр.
# shortcut: пороги подобраны на 9 колодах; при новых жалобах на пропуск — добавить колоду в набор и перепроверить
GRID, EMPTY, CARD_GRID, CARD_EMPTY = 16, 0.30, 8, 0.40
EDGE_STEP, BUSY = (
    24,
    0.06,
)  # перепад яркости на пиксель и доля таких пикселей, с которой зона — картинка
BACKGROUND, CONTAINER, MIN_CARD, EDGE = 0.90, 0.25, 0.03, 0.08
PT, INSET, DEFAULT_PT = 12700, 91440, 18
NUM = re.compile(r"^\s*\d{1,2}\s*/\s*\d{1,2}\s*$")
# рабочие пометки агента, попавшие на слайд
WORKNOTE = re.compile(
    r"\bне\s+найден[аоы]?\b|первоисточник\s+не|\bTODO\b|\bTBD\b|\bFIXME\b|\?\?\?|\bуточнить\b|заглушк|\bplaceholder\b|Рамка:\s*ТЗ",
    re.I,
)
# служебная строка титула: курс, тема, вид работы, ссылка на главу или источник
# калибровка 27.10.2026: 1063 строки первых слайдов 222 учебных колод — 90 служебных находятся;
# темы «Источники экономического роста», «Глава 25 НК РФ: …», «Санкции — доклад МВФ» не срабатывают
SEP = r"[·•|—–]"
DOT = r"[·•|]"  # «часть II» после тире — часть темы («Налоговый кодекс — часть II»)
KW = (
    r"(?:тема\s*№?\s*\d+|семинар\w*|лекци[яи]\s*№?\s*\d*|заняти[ея]\s*№?\s*\d*|модуль\s*№?\s*\d+"
    r"|доклад\w*|реферат\w*|дз\s*-?\s*\d+|курсов\w*\s+работ\w*)"
)
PART = r"часть\s+(?:\d+|[IVXLC]+)\b"
TAIL = rf"\s*(?:(?:доклад\w*|работ\w*)\s*)?(?=$|{SEP})"  # после вида работы — конец строки или разделитель
NUMBERED = (
    r"(?:тема|семинар|лекци[яи]|заняти[ея]|модуль|часть|дз)\s*[-№]?\s*(?:\d+|[IVXLC]+)"
)
TITLE = re.compile(
    rf"{SEP}\s*{KW}{TAIL}|{DOT}\s*{PART}{TAIL}|\b{NUMBERED}\s*{SEP}|^\s*{NUMBERED}\s*$"
    r"|^\s*(?:учебный|тезисный)\s+доклад|^\s*(?:доклад|реферат)(?:\s+по\b|\s*$)|домашн\w*\s+задани\w*"
    r"|^\s*по\s+(?:материалам\s+)?глав|^\s*обзор\s+стат(?:ей|ьи)\s*:|\bс\.\s*\d+|^\s*источник(?:и)?\s*:"
    r"|\bглава\s+\d+\s*(?:[·•|—–]|\.(?!\d)|$)",  # «Глава 26. …», «Бланшар, глава 26»; не «Глава 25 НК РФ: …», «Глава 26.2 НК РФ»
    re.I | re.M,
)


def title_hits(body):
    """Служебные строки титула; мягкий перенос строки (\x0b у python-pptx) — тоже конец строки."""
    return [m.group(0) for m in TITLE.finditer(body.replace("\x0b", "\n"))]


def yo(text):
    return text.replace("ё", "е").replace("Ё", "Е")


def author_rx(name):
    """Фамилия в падежах, с заглавной или капсом, текст — через yo(): «Кузнец» — не «Кузнецов», «Ли» — не «лишь», «Смирнова» — и «Смирновой»."""
    name = yo(name.strip())
    if name.isupper() or name.islower():  # «КУЗНЕЦ», «кузнец» → «Кузнец»; «МакКартни», «Де Голль» — как даны
        name = re.sub(r"[^\s-]+", lambda m: m.group(0)[:1].upper() + m.group(0)[1:].lower(), name)
    low = name.lower()
    if len(name) <= 2 or low[-1] in "оеиуюыэ":  # двухбуквенные и несклоняемые — только как есть
        stem, ends = name, ("",)
    elif low.endswith("ая"):
        stem, ends = name[:-2], ("ая", "ой", "ую")
    elif low.endswith(("ий", "ый", "ой")) and len(name) > 3:  # Толстой склоняется как прилагательное, Цой — нет
        stem, ends = name[:-2], (low[-2:], "ого", "ому", "им", "ым", "ом")
    elif low[-1] in "ая":
        stem, ends = name[:-1], (low[-1], "ы", "и", "е", "у", "ю", "ой", "ей", "ою")
    elif low[-1] in "ьй":
        stem, ends = name[:-1], (low[-1], "я", "ю", "ем", "е")
    else:
        stem, ends = name, ("", "а", "у", "ом", "ем", "ым", "е")
    stem, forms = stem[:1].upper() + stem[1:], "|".join(ends)
    return re.compile(
        rf"\b(?:{re.escape(stem)}(?:{forms})|{re.escape(stem.upper())}(?:{forms.upper()}))\b"
    )


SHAPES = (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM)


def leaves(shapes, tf=(0.0, 0.0, 1.0, 1.0)):
    """(фигура, преобразование в EMU слайда X = ox + x·sx) с учётом off/ext/chOff/chExt групп."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            x = sh._element.grpSpPr.find(qn("a:xfrm"))
            if x is None:
                yield from leaves(sh.shapes, tf)
                continue
            num = lambda tag, attr: (
                int(x.find(qn(tag)).get(attr)) if x.find(qn(tag)) is not None else 0
            )  # noqa: E731
            gx, gy, gw, gh = (
                num("a:off", "x"),
                num("a:off", "y"),
                num("a:ext", "cx"),
                num("a:ext", "cy"),
            )
            cx, cy = num("a:chOff", "x"), num("a:chOff", "y")
            cw, ch = num("a:chExt", "cx") or gw or 1, num("a:chExt", "cy") or gh or 1
            ox, oy, sx, sy = tf
            kx, ky = gw / cw if gw else 1, gh / ch if gh else 1
            yield from leaves(
                sh.shapes,
                (ox + sx * (gx - cx * kx), oy + sy * (gy - cy * ky), sx * kx, sy * ky),
            )
        else:
            yield sh, tf


def frame(sh, tf):
    if None in (sh.left, sh.top, sh.width, sh.height):
        return None
    ox, oy, sx, sy = tf
    return (
        ox + sh.left * sx,
        oy + sh.top * sy,
        ox + (sh.left + sh.width) * sx,
        oy + (sh.top + sh.height) * sy,
    )


def text_box(sh, box):
    """Площадь, которую реально занимает текст внутри рамки (EMU)."""
    tf = sh.text_frame
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None and body.get("wrap") == "none":
        return box  # без переноса рамка облегает текст; spAutoFit с переносом бывает выше текста
    norm = body.find(qn("a:normAutofit")) if body is not None else None
    scale = int(norm.get("fontScale", 100000)) / 100000 if norm is not None else 1
    x0, y0, x1, y1 = box
    width = max(x1 - x0 - 2 * INSET, PT)
    used_w, used_h = 0.0, 2 * INSET
    for p in tf.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size] or [DEFAULT_PT]
        size = max(sizes) * scale * PT
        line = len(p.text) * 0.55 * size
        used_w = max(used_w, min(line, width))
        used_h += max(1, math.ceil(line / width)) * 1.2 * size
    used_w += 2 * INSET
    anchor = body.get("anchor", "t") if body is not None else "t"
    top = y0 + {"ctr": (y1 - y0 - used_h) / 2, "b": y1 - y0 - used_h}.get(anchor, 0)
    align = str(tf.paragraphs[0].alignment or "")
    left = x0 + (
        (x1 - x0 - used_w) / 2
        if "CENTER" in align
        else (x1 - x0 - used_w)
        if "RIGHT" in align
        else 0
    )
    return max(x0, left), max(y0, top), min(x1, left + used_w), min(y1, top + used_h)


def classify(slide, sw, sh_):
    """Содержимое и контейнеры слайда в долях слайда: (content, containers, texts)."""
    content, containers, maybe, texts = [], [], [], []
    norm = lambda b: (
        max(0, b[0] / sw),
        max(0, b[1] / sh_),
        min(1, b[2] / sw),
        min(1, b[3] / sh_),
    )  # noqa: E731
    for sh, tf in leaves(slide.shapes):
        box = frame(sh, tf)
        if box is None:
            continue
        nb = norm(box)
        area = max(0, nb[2] - nb[0]) * max(0, nb[3] - nb[1])
        has_text = getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()
        if has_text:
            texts.append((sh, nb))
            content.append(norm(text_box(sh, box)))
            if sh.shape_type in SHAPES and area >= MIN_CARD:
                containers.append(nb)
        elif getattr(sh, "has_chart", False) or getattr(sh, "has_table", False):
            content.append(nb)
            if sh.has_table:
                texts.append((sh, nb))
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
            if area < BACKGROUND:
                content.append(nb)
        elif sh.shape_type in SHAPES:
            if area >= BACKGROUND:
                continue
            if area >= CONTAINER:
                containers.append(nb)
            elif area >= MIN_CARD:
                maybe.append(nb)
            else:
                content.append(nb)
    for nb in maybe:  # фигура с блоками внутри — карточка, без них — элемент оформления
        (containers if any(inside(c, nb) for c in content) else content).append(nb)
    return content, containers, texts


def overlaps(box, outer):
    return (
        box[0] < outer[2]
        and box[2] > outer[0]
        and box[1] < outer[3]
        and box[3] > outer[1]
    )


def clip(box, outer):
    return (
        max(box[0], outer[0]),
        max(box[1], outer[1]),
        min(box[2], outer[2]),
        min(box[3], outer[3]),
    )


def inside(box, outer):
    cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
    return outer[0] <= cx <= outer[2] and outer[1] <= cy <= outer[3]


def empty_share(boxes, area=(0.0, 0.0, 1.0, 1.0), grid=GRID):
    """Доля самого большого прямоугольника сетки внутри area, не задетого ни одним box."""
    return largest_empty(boxes, area, grid)[0]


def largest_empty(boxes, area=(0.0, 0.0, 1.0, 1.0), grid=GRID):
    """(доля, прямоугольник в долях слайда) самого большого пустого прямоугольника сетки."""
    ax0, ay0, ax1, ay1 = area
    cw, ch = (ax1 - ax0) / grid, (ay1 - ay0) / grid
    busy = [[False] * grid for _ in range(grid)]
    for x0, y0, x1, y1 in boxes:
        for r in range(grid):
            for c in range(grid):
                cx0, cy0 = ax0 + c * cw, ay0 + r * ch
                if x0 < cx0 + cw and x1 > cx0 and y0 < cy0 + ch and y1 > cy0:
                    busy[r][c] = True
    best, rect, heights = 0, area, [0] * grid
    for r, row in enumerate(busy):
        heights = [0 if b else h + 1 for b, h in zip(row, heights)]
        stack = []
        for i, h in enumerate(heights + [0]):
            start = i
            while stack and stack[-1][1] >= h:
                start, top = stack.pop()
                if top * (i - start) > best:
                    best = top * (i - start)
                    rect = (
                        ax0 + start * cw,
                        ay0 + (r - top + 1) * ch,
                        ax0 + i * cw,
                        ay0 + (r + 1) * ch,
                    )
            stack.append((start, h))
    return best / (grid * grid), rect


def busy_picture(png, rect):
    """Доля резких перепадов яркости в зоне PNG рендера: высокая — там картинка, а не пустота."""
    import numpy as np
    from PIL import Image

    img = Image.open(png).convert("L")
    w, h = img.size
    crop = np.asarray(
        img.crop(
            (int(rect[0] * w), int(rect[1] * h), int(rect[2] * w), int(rect[3] * h))
        ),
        dtype=float,
    )
    if crop.size < 4:
        return 0.0
    grad = np.abs(np.diff(crop, axis=0))[:, :-1] + np.abs(np.diff(crop, axis=1))[:-1, :]
    return float((grad > EDGE_STEP).mean())


def content_boxes(slide, sw, sh_):
    return classify(slide, sw, sh_)[0]


def check(
    deck_path,
    *,
    skip=None,
    accept=(),
    empty=EMPTY,
    strict=False,
    allow=None,
    render=None,
    authors=(),
    allow_title=False,
):
    prs = Presentation(deck_path)
    authors = [(name, author_rx(name)) for name in authors]
    sw, sh_ = prs.slide_width, prs.slide_height
    n = len(prs.slides)
    skip = {1, n} if skip is None else set(skip)
    result = {}
    pngs = sorted(Path(render).glob("slide-*.png")) if render else []
    shown = sum(s._element.get("show") != "0" for s in prs.slides)
    if pngs and len(pngs) not in (
        n,
        shown,
    ):  # рендер не от этой колоды: зону сверять не с чем
        raise ValueError(f"в {render} PNG {len(pngs)}, слайдов {n} (видимых {shown})")
    by_slide = (
        len(pngs) == n
    )  # рендер со скрытыми слайдами (Windows) — по номеру слайда
    visible = 0
    for no, slide in enumerate(prs.slides, 1):
        if slide._element.get("show") == "0":
            continue
        visible += 1
        found = []
        add = lambda kind, level, detail: found.append(
            {"kind": kind, "level": level, "detail": detail}
        )  # noqa: E731
        void_level = "accepted" if no in accept else "error"
        content, containers, texts = classify(slide, sw, sh_)
        if no not in skip:
            share, rect = largest_empty(content)
            png = (pngs[no - 1] if by_slide else pngs[visible - 1]) if pngs else None
            if share >= empty and png and busy_picture(png, rect) >= BUSY:
                pass  # в «пустой» зоне на рендере картинка во весь слайд
            elif share >= empty:
                add("void", void_level, f"пустая зона {share:.0%} слайда")
            for c in containers:
                inner = [clip(b, c) for b in content if overlaps(b, c)]
                if inner and (share := empty_share(inner, c, CARD_GRID)) >= CARD_EMPTY:
                    add("card", void_level, f"пустая часть карточки {share:.0%}")
        for sh, nb in texts:
            if (
                sh.is_placeholder
                and sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
            ):
                add("counter", "error", "номер слайда")
            body = (
                "\n".join(c.text for row in sh.table.rows for c in row.cells)
                if getattr(sh, "has_table", False)
                else sh.text_frame.text
            )
            if (
                not getattr(sh, "has_table", False)
                and NUM.match(body)
                and (nb[1] < EDGE or nb[3] > 1 - EDGE)
            ):
                add("counter", "error", f"счётчик «{body.strip()}»")
            for m in WORKNOTE.finditer(body):
                add("worknote", "error", f"рабочая пометка «{m.group(0)}»")
            if no == 1 and not allow_title:
                for hit in title_hits(body):
                    add("title", "error", f"служебная строка титула «{hit}»")
            for name, rx in authors if no > 1 else ():
                if rx.search(yo(body)):
                    add("authors", "error", f"фамилия автора «{name}» вне титула")
            for frag, _ in stoplist.findings(body, strict, allow):
                add("stop", "error" if strict else "warn", f"«{frag}»")
        if slide.has_notes_slide:
            for frag, _ in stoplist.findings(
                slide.notes_slide.notes_text_frame.text, strict, allow
            ):
                add("stop", "error" if strict else "warn", f"заметки: «{frag}»")
        if found:
            result[no] = found
    return result


def exit_code(result):
    return int(any(f["level"] == "error" for fs in result.values() for f in fs))


def selftest():
    assert empty_share([]) == 1.0
    assert empty_share([(0, 0, 1, 1)]) == 0.0
    assert abs(empty_share([(0, 0, 0.5, 1)]) - 0.5) < 1e-9
    assert NUM.match("03/12") and not NUM.match("24/7 поддержка")
    assert WORKNOTE.search("Спред к своп-ставке не найден") and WORKNOTE.search(
        "TODO: цифра"
    )
    assert not WORKNOTE.search("найденные решения") and not WORKNOTE.search(
        "Источник: INSEE, 2025"
    )
    for kicker in (
        "МИКРОЭКОНОМИКА · ТЕМА 5",
        "МАКРОЭКОНОМИКА · СЕМИНАР",
        "УЧЕБНЫЙ ДОКЛАД",
        "По главе 26 учебника",
        "Маркетинг · Часть III",
        "Бланшар, с. 585",
        "МИРОВАЯ ЭКОНОМИКА II · ДОКЛАД",
        "ФИНАНСЫ И КРЕДИТ · РЕФЕРАТ",
        "Экономика | ДЗ-4",
        "Обзор статей: Barnichon, Matthes, Price (2017)",
        "Источники: [1]",
        "Глава 26. Долг и дефицит",
        "Глава 9 · Деньги",
        "ГЛАВА 12",
        "Бланшар, глава 26",
    ):
        assert TITLE.search(kicker), kicker
    assert title_hits("МИРОВАЯ ЭКОНОМИКА II · ДОКЛАД\x0bМиграционная политика ЕС")
    assert title_hits("МИКРОЭКОНОМИКА · ТЕМА 5\x0bРынок платформ")
    for line in (
        "ВВП · темпы роста",
        "Кузнец Иван · Петров Олег",
        "Обзор статистики безработицы",
        "Алиев Тимур, Кузнец Иван 402 группа",
        "Экспорт · модульные дома",
        "Источники экономического роста",
        "Обзор статей расходов федерального бюджета",
        "Налог на прибыль (глава 25 НК РФ)",
        "Глава 25 НК РФ: налог на прибыль",
        "Налоговый кодекс — часть II",
        "Санкции — доклад МВФ",
        "Бюджет 2026 | часть 1: доходы",
        "Пенсионная реформа — лекция Кудрина",
        "Глава 26.2 НК РФ: упрощённая система налогообложения",
        "Глава 26.1 НК РФ: ЕСХН",
        "Глава 25.3. Государственная пошлина",
    ):
        assert not TITLE.search(line), line
    smith, li = author_rx("Кузнец"), author_rx("Ли")
    assert smith.search("доклад Кузнеца") and smith.search("КУЗНЕЦ И.")
    assert not any(smith.search(s) for s in ("Кузнецов", "Кузнечик", "кузнецы"))
    assert li.search("Ли Куан Ю") and not li.search("Лишь лишь")
    assert not author_rx("Пак").search("пакет") and author_rx("Шевченко").search(
        "Шевченко"
    )
    assert author_rx("Смирнова").search("у Смирновой") and author_rx("Алиев").search(
        "с Алиевым"
    )
    assert not author_rx("Алиев").search("Алиевский")
    assert author_rx("Пак").search("у Пака") and author_rx("Ким").search("с Кимом") and not author_rx("Ким").search("Кимберли")
    assert author_rx("Цой").search("песни Цоя") and author_rx("Толстой").search("романы Толстого")
    assert author_rx("КУЗНЕЦ").search("Кузнец") and author_rx("Королёв").search(yo("у Королёва"))
    assert author_rx("МакКартни").search("песни МакКартни") and author_rx("Де Голль").search("политика Де Голля")
    assert author_rx("Кузнец Иван").search("доклад: Кузнец Иван") and author_rx("ДЕ ГОЛЛЬ").search("Де Голля")
    print("selftest ok")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck", nargs="?")
    ap.add_argument("--skip", help="слайды без проверки пустот, через запятую")
    ap.add_argument(
        "--accept", default="", help="слайды с намеренной пустотой, через запятую"
    )
    ap.add_argument("--empty", type=float, default=EMPTY)
    ap.add_argument(
        "--strict",
        action="store_true",
        help="учебная сдача: стоп-фразы и личные пометки — ошибки",
    )
    ap.add_argument(
        "--allow", help="регулярка: разрешённый контекст стоп-фразы (колода про ИИ)"
    )
    ap.add_argument(
        "--render",
        help="папка PNG рендера (5-проверка/roundN): зона с картинкой во весь слайд — не пустота",
    )
    ap.add_argument(
        "--authors",
        default="",
        help="фамилии авторов через запятую: вне титула — находка",
    )
    ap.add_argument(
        "--allow-title",
        action="store_true",
        help="задание требует курс или вид работы на титуле (раздел «Требования к сдаче» памяти предмета)",
    )
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if not a.deck:
        ap.error("нужен deck.pptx")
    nums = lambda s: {int(x) for x in s.split(",") if x.strip()}  # noqa: E731
    try:
        result = check(
            a.deck,
            skip=nums(a.skip) if a.skip is not None else None,
            accept=nums(a.accept),
            empty=a.empty,
            strict=a.strict,
            allow=a.allow,
            render=a.render,
            authors=[x.strip() for x in a.authors.split(",") if x.strip()],
            allow_title=a.allow_title,
        )
    except ValueError as e:
        sys.exit(f"ошибка: {e}")
    if a.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for no, fs in result.items():
            for f in fs:
                print(f"слайд {no}: [{f['level']}] {f['detail']}")
        print("находок нет" if not result else f"слайдов с находками: {len(result)}")
    sys.exit(exit_code(result))


if __name__ == "__main__":
    main()
