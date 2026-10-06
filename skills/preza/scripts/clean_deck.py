#!/usr/bin/env python3
"""Чистая версия колоды: без переходов, анимаций и промежуточных слайдов Morph.

    python3 clean_deck.py <показ.pptx> <svg_output> -o <чистая.pptx>
    python3 clean_deck.py --selftest

Промежуточный слайд — страница svg_output с буквой после номера (`10b_bfire-pan.svg`).
Слайды сопоставляются со страницами по порядку экспорта ppt-master: числа в имени по значению.
"""

import argparse
import io
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
STAGING = re.compile(r"^\d+[a-z]", re.IGNORECASE)


def order_key(name):
    """Та же сортировка, что slide_roster.discover_slide_svgs в ppt-master."""
    folded = name.casefold()
    return tuple(
        (0, int(s)) if s.isdigit() else (1, s) for s in re.split(r"(\d+)", folded)
    ), name


def strip_motion(sld):
    """Снять переход и анимации со слайда (включая переходы внутри mc:AlternateContent)."""
    removed = 0
    for el in list(sld):
        if el.tag in (P + "transition", P + "timing") or (
            el.tag == MC + "AlternateContent"
            and el.find(f".//{P}transition") is not None
        ):
            sld.remove(el)
            removed += 1
    return removed


def clean(prs, svg_names):
    names = sorted(svg_names, key=order_key)
    if len(prs.slides) != len(names):
        raise ValueError(
            f"слайдов {len(prs.slides)}, страниц {len(names)}: колода и svg_output разошлись — переэкспортируй показ"
        )
    staging = [i for i, n in enumerate(names) if STAGING.match(n)]
    motion = sum(strip_motion(s._element) for s in prs.slides)
    ids = prs.slides._sldIdLst
    for i in reversed(staging):
        sld_id = ids[i]
        prs.part.drop_rel(sld_id.rId)
        ids.remove(sld_id)
    return {
        "убрано промежуточных": [names[i] for i in staging],
        "снято элементов движения": motion,
        "осталось слайдов": len(prs.slides),
    }


def selftest():
    prs = Presentation()
    for text in ("A", "B", "B-pan", "C"):
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            0, 0, 100, 100
        ).text_frame.text = text
    ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    prs.slides[1]._element.append(
        parse_xml(f"<p:transition {ns}><p:fade/></p:transition>")
    )
    prs.slides[2]._element.append(
        parse_xml(
            f'<p:timing {ns}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot"/></p:par></p:tnLst></p:timing>'
        )
    )
    names = [
        "03_c.svg",
        "01_a.svg",
        "02b_b-pan.svg",
        "02_b.svg",
    ]  # порядок на диске не важен
    report = clean(prs, names)
    assert report["убрано промежуточных"] == ["02b_b-pan.svg"], report
    assert report["снято элементов движения"] == 2, report
    buf = io.BytesIO()
    prs.save(buf)
    back = Presentation(io.BytesIO(buf.getvalue()))
    assert [s.shapes[0].text_frame.text for s in back.slides] == ["A", "B", "C"]
    assert not any(
        el.tag in (P + "transition", P + "timing")
        for s in back.slides
        for el in s._element
    )
    try:
        clean(back, names)
    except ValueError:
        pass
    else:
        raise AssertionError(
            "расхождение слайдов и страниц должно останавливать скрипт"
        )
    assert sorted(["10b_x.svg", "10_x.svg", "9_x.svg"], key=order_key) == [
        "9_x.svg",
        "10_x.svg",
        "10b_x.svg",
    ]
    print("selftest ok")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck", nargs="?")
    ap.add_argument("svg_output", nargs="?")
    ap.add_argument("-o", "--output")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if not (a.deck and a.svg_output and a.output):
        ap.error("нужны <показ.pptx> <svg_output> -o <чистая.pptx>")
    prs = Presentation(a.deck)
    try:
        report = clean(prs, [p.name for p in Path(a.svg_output).glob("*.svg")])
    except ValueError as e:
        sys.exit(str(e))
    prs.save(a.output)
    for k, v in report.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
