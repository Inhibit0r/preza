"""Проверка происхождения чисел по реестру утверждений.

Реестр: data/claims.csv (значение, период, источник, контекст), применение по слайдам: data/usage.csv.
- claims: calc-источники сверяются с data/metrics.json, остальные id — с data/sources.md;
- --script script.md: каждое число в предложении слайда N должно совпасть (со знаком) с утверждением
  из usage[N], и в том же предложении должно быть слово-контекст этого утверждения;
- --pptx deck.pptx: текст фигур и таблиц, категории и значения рядов нативных графиков слайда N
  сверяются с утверждениями usage[N]; сопоставление «число → утверждение» пишется в отчёт
  для содержательной сверки при аудите.
Годы (четыре цифры подряд, 19xx–20xx) и номера «Тема/Слайд/Эпизод N» не проверяются. Слайд с usage = SKIP (источники) пропускается.

Скрипт лежит в <проект>/data/ (копия из навыка preza). Запуск: python3 data/check_numbers.py [--script script.md] [--pptx deck.pptx --report out.md]
"""

import argparse
import csv
import json
import re
from pathlib import Path

root = Path(__file__).resolve().parents[1]
metrics_file = root / "data/metrics.json"  # нет расчётов — нет файла
metrics = json.loads(metrics_file.read_text()) if metrics_file.exists() else {}
sources = (root / "data/sources.md").read_text()
claims = {
    c["id"]: c for c in csv.DictReader((root / "data/claims.csv").open(), delimiter=";")
}
usage = {
    int(r["slide"]): r["claims"]
    for r in csv.DictReader((root / "data/usage.csv").open(), delimiter=";")
}

DATE = re.compile(
    r"(?<![\d.])(\d{2}\.)?(\d{2})\.(\d{4})(?!\.?\d)"
)  # 02.12.2024 и 12.2024
NUM = re.compile(
    r"(?<![\d.,A-Za-zА-Яа-яЁё])([−+-])?(\d{1,3}(?:[   ]\d{3})+|\d+)(?:[.,](\d+))?"
)
DECLINE = re.compile(r"сниз|сниж|упал|паден|минус|сократ|подешев|дешевле", re.I)
GROWTH = re.compile(r"вырос|рост|увелич|подорож|\+", re.I)
SKIP_LABELS = re.compile(
    r"(Тем[аеуы]|Слайд|Слайды|№|ЭПИЗОД|Эпизод)\s*\d+|\b[1-4]К\b"
)  # «1К 2018» — квартал
errors = []


def to_float(s):
    return float(re.sub(r"[   ]", "", s).replace(",", ".").replace("−", "-"))


def claim_dates(c):
    return {m.group(0) for m in DATE.finditer(c["value"] + " " + c["period"])}


# --- 1. Реестр: calc-значения и источники ---
for c in claims.values():
    for src in re.split(r"\s*/\s*", c["source"]):
        if src.startswith("calc:"):
            key, *sub = src[5:].split(".", 1)
            got = metrics[key][sub[0]] if sub else metrics[key]
            decimals = len(c["value"].split(",")[1]) if "," in c["value"] else 0
            if round(float(got), decimals) != to_float(c["value"]):
                errors.append(
                    f"{c['id']}: claims {c['value']} ≠ metrics {src[5:]}={got}"
                )
        elif f"| {src} |" not in sources:
            errors.append(f"{c['id']}: источник {src} не найден в sources.md")
for slide, ids in usage.items():
    if ids != "SKIP":
        errors += [
            f"usage слайд {slide}: нет утверждения {i}"
            for i in ids.split(",")
            if i not in claims
        ]


def check_fragment(where, slide, text, context, require_context, report):
    """Проверить все числа фрагмента text; context — текст, где ищутся слова-контексты."""
    if usage.get(slide) == "SKIP":
        return
    allowed = [claims[i] for i in usage.get(slide, "").split(",") if i in claims]
    text = SKIP_LABELS.sub(" ", text)
    dates = set().union(*(claim_dates(c) for c in allowed)) if allowed else set()
    for m in DATE.finditer(text):
        if m.group(0) not in dates:
            errors.append(
                f"{where}: дата {m.group(0)} не относится к утверждениям слайда {slide}"
            )
    text = DATE.sub(" ", text)
    for m in NUM.finditer(text):
        sign, whole, frac = m.groups()
        prev = text[max(0, m.start() - 1) : m.start()]
        explicit = sign if sign and not prev.isdigit() else ""
        val = to_float(whole + ("." + frac if frac else ""))
        if not frac and re.fullmatch(r"(19|20)\d\d", whole):  # год, а не «2 000»
            continue
        token = m.group(0).strip("  ")
        hits = []
        for c in allowed:
            cv = to_float(c["value"]) if not DATE.fullmatch(c["value"]) else None
            if cv is None or abs(cv) != val:
                continue
            if explicit in ("−", "-") and cv >= 0 or explicit == "+" and cv < 0:
                continue
            if (
                not explicit
                and cv < 0
                and not (DECLINE.search(context) and not GROWTH.search(context))
            ):
                continue
            if require_context and not re.search(c["context"], context, re.I):
                continue
            hits.append(c["id"])
        if hits:
            report.append(
                f"| {slide} | {token} | {', '.join(hits)} | {claims[hits[0]]['claim']} |"
            )
        else:
            errors.append(
                f"{where}: число {token} не привязано к утверждению слайда {slide}"
            )


def check_script(path, report):
    text = Path(path).read_text()
    parts = re.split(r"^## Слайд (\d+)\..*$", text, flags=re.M)
    for i in range(1, len(parts), 2):
        slide = int(parts[i])
        for sent in re.split(r"(?<=[.!?])\s+", parts[i + 1]):
            check_fragment(f"{path} слайд {slide}", slide, sent, sent, True, report)


def slide_fragments(shapes):
    """Весь видимый текст слайда: фигуры (в т.ч. внутри групп), ячейки таблиц, графики."""
    frags = []
    for shape in shapes:
        if hasattr(shape, "shapes"):  # группа
            frags += slide_fragments(shape.shapes)
        if shape.has_text_frame:
            frags.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            frags += [cell.text for row in shape.table.rows for cell in row.cells]
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = shape.chart
            if chart.has_title:
                frags.append(chart.chart_title.text_frame.text)
            for axis in ("category_axis", "value_axis"):
                try:
                    ax = getattr(chart, axis)
                    if ax.has_title:
                        frags.append(ax.axis_title.text_frame.text)
                except (ValueError, NotImplementedError):
                    pass
            for plot in chart.plots:
                frags += [str(cat) for cat in plot.categories]
                for series in plot.series:
                    frags.append(series.name or "")
                    frags += [
                        f"{v:g}".replace(".", ",")
                        for v in series.values
                        if v is not None
                    ]
    return frags


def check_pptx(path, report):
    from pptx import Presentation

    for n, slide in enumerate(Presentation(path).slides, start=1):
        frags = slide_fragments(slide.shapes)
        whole = "\n".join(frags)
        # ponytail: на слайде контекст = весь текст слайда; подписи без контекста сверяет аудит Codex по отчёту
        for frag in frags:
            check_fragment(f"{path} слайд {n}", n, frag, whole, False, report)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--script")
    ap.add_argument("--pptx")
    ap.add_argument("--report")
    a = ap.parse_args()
    report = ["| слайд | число | утверждение | формулировка |", "|---|---|---|---|"]
    if a.script:
        check_script(a.script, report)
    if a.pptx:
        check_pptx(a.pptx, report)
    if a.report:
        Path(a.report).parent.mkdir(parents=True, exist_ok=True)
        Path(a.report).write_text("\n".join(report) + "\n")
    print(
        "claims:",
        len(claims),
        "| чисел сопоставлено:",
        len(report) - 2,
        "| ошибок:",
        len(errors),
    )
    for e in errors:
        print("  -", e)
    assert not errors
