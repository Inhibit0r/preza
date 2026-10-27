#!/usr/bin/env python3
"""Проверка речи: разделы «## Слайд N. …», объём под хронометраж, стоп-фразы, заметки PPTX, Word.

  python3 check_speech.py script.md --slides 13 [--words 720-860] [--forbid REGEX] [--pptx deck.pptx]
                          [--docx speech.docx] [--strict] [--allow REGEX]

--words — цель ±10 % (минуты × 130 слов, при плотных числах × 120): короче — ошибка, длиннее — предупреждение
(цель мягкая: один проход, убрать явно лишнее; лишнего нет — оставить). Без --words объём не проверяется:
так сверяют сданный Word со снимком сданной речи (script.delivered.md --docx).

Стоп-список сдачи (stoplist.py) ищется в речи, в слайдах (--pptx) и в Word (--docx): предупреждение,
при --strict (учебная сдача) — ошибка. --forbid — дополнительные запреты пользователя (всегда ошибка).
--pptx сверяет заметки докладчика каждого слайда с разделом речи (пробелы не учитываются).
--docx: шапка, заголовки и текст разделов «Слайд N.» в Word совпадают с речью (без учёта пробелов, «ё»/«е», вида тире, разметки md) —
иначе ошибка: ловит и ручные правки одного из файлов; «—» и «ё» в Word — предупреждение, при --strict — ошибка.
Самопроверка: --selftest
"""

import argparse
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

sys.path.insert(0, str(Path(__file__).resolve().parent))
import stoplist  # noqa: E402

HEAD = re.compile(r"^## Слайд \d+\..*$", re.M)
# хронометраж в скобках в заголовке слайда: «(1 мин)», «(~40 сек)», «(0:45)»
TIMED = re.compile(
    r"^(?:## )?Слайд \d+\..*\((?=[^)]*\d)[^)]*(?:мин|сек|\d:\d\d)[^)]*\)", re.M
)


def words(text):
    return len(re.findall(r"[\wЁё-]+", text))


def shape_texts(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):  # группа
            yield from shape_texts(shape.shapes)
        if shape.has_text_frame:
            yield shape.text_frame.text
        if getattr(shape, "has_table", False) and shape.has_table:
            yield from (cell.text for row in shape.table.rows for cell in row.cells)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WORD_HEAD = re.compile(r"^Слайд \d+\..*$", re.M)


def docx_text(path):
    """Текст Word по абзацам из word/document.xml (stdlib: python-docx не нужен)."""
    with zipfile.ZipFile(path) as z:
        root = ElementTree.fromstring(z.read("word/document.xml"))
    return "\n".join(
        "".join(t.text or "" for t in p.iter(W + "t")) for p in root.iter(W + "p")
    )


def norm(text):
    text = text.replace("ё", "е").replace("Ё", "Е").replace("\u00ad", "")
    return " ".join(re.sub(r"[—–‑]", "-", text).split())


def split_sections(text, head):
    """Шапка до первого раздела и [(заголовок, тело)]."""
    ms = list(head.finditer(text))
    ends = [m.start() for m in ms[1:]] + [len(text)]
    return text[: ms[0].start()] if ms else text, [
        (m.group(0), text[m.end() : end]) for m, end in zip(ms, ends)
    ]


def plain(text):
    return norm(re.sub(r"^#+\s*|\*\*|__", "", text, flags=re.M))


def word_mismatch(text, body):
    """Где Word расходится с речью — шапка, заголовки, разделы: Word собирают из script.md, а правят часто одно из двух."""
    head_s, secs = split_sections(text, HEAD)
    head_w, parts = split_sections(body, WORD_HEAD)
    if len(parts) != len(secs):
        return [f"Word: разделов «Слайд N.» {len(parts)}, в речи {len(secs)}"]
    found = (
        ["Word: шапка не совпадает с речью"] if plain(head_s) != plain(head_w) else []
    )
    for n, ((hs, bs), (hw, bw)) in enumerate(zip(secs, parts), 1):
        if plain(hs) != plain(hw):
            found.append(f"Word: заголовок слайда {n} не совпадает с речью")
        if plain(bs) != plain(bw):
            found.append(f"Word: слайд {n} не совпадает с речью")
    return found


def words_range(value):
    """«720-860» или цель «790» (±10 %)."""
    if "-" in value:
        lo, hi = (int(x) for x in value.split("-", 1))
    else:
        lo, hi = round(int(value) * 0.9), round(int(value) * 1.1)
    return lo, hi


def check(
    text, slides, lo, hi, forbid=None, deck=None, *, strict=False, allow=None, word=None
):
    """(ошибки, предупреждения, слов по разделам); пустой список ошибок — речь в порядке."""
    secs = HEAD.split(text)[1:]
    per = [words(s) for s in secs]
    errors, warnings = [], []
    soft = errors if strict else warnings
    if len(secs) != slides:
        errors.append(f"разделов {len(secs)}, слайдов {slides}")
    if lo is not None and sum(per) < lo:
        errors.append(f"слов {sum(per)}, нужно {lo}–{hi}")
    elif hi is not None and sum(per) > hi:
        warnings.append(
            f"слов {sum(per)} — длиннее цели {lo}–{hi}: один проход, убрать явно лишнее; лишнего нет — оставить"
        )
    texts = [text]
    if deck is not None:
        if len(deck.slides) != slides:
            errors.append(f"в PPTX {len(deck.slides)} слайдов, ожидалось {slides}")
        for n, (slide, sec) in enumerate(zip(deck.slides, secs), 1):
            notes = (
                slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            )
            if " ".join(notes.split()) != " ".join(sec.split()):
                errors.append(f"слайд {n}: заметки докладчика не совпадают с речью")
            texts += shape_texts(slide.shapes)
    if word is not None:
        body = docx_text(word)
        texts.append(body)
        errors += word_mismatch(text, body)
        for sign, name in (("—", "длинное тире"), ("ё", "«ё»")):
            if count := body.lower().count(sign):
                soft.append(f"Word: {name} — {count}")
    timed = [
        m.group(0)
        for t in (text, body if word is not None else "")
        for m in TIMED.finditer(t)
    ]
    if timed:
        soft.append(f"время в заголовках слайдов: {timed}")
    if forbid:
        found = [m.group(0) for t in texts for m in re.finditer(forbid, t, re.I)]
        if found:
            errors.append(f"запрещённые формулировки: {found}")
    stop = [frag for t in texts for frag, _ in stoplist.findings(t, strict, allow)]
    if stop:
        soft.append(f"стоп-фразы сдачи: {stop}")
    return errors, warnings, per


def selftest():
    speech = "# Текст выступления · около 1 минуты · 2 слайда\nИванов И.\n\n## Слайд 1. А\nЕщё — раз.\n## Слайд 2. Б\nДва.\n"
    word = "Текст выступления · около 1 минуты · 2 слайда\nИванов И.\nСлайд 1. А\nЕще - раз.\nСлайд 2. Б\nДва."
    assert word_mismatch(speech, word) == []
    assert word_mismatch(speech, word.replace("Два.", "Два, три.")) == [
        "Word: слайд 2 не совпадает с речью"
    ]
    assert word_mismatch(speech, word.replace("Иванов И.", "Иванов Иван, 402")) == [
        "Word: шапка не совпадает с речью"
    ]  # правка пользователя в шапке Word тоже видна
    assert word_mismatch(speech, word.replace("Слайд 2. Б", "Слайд 2. В")) == [
        "Word: заголовок слайда 2 не совпадает с речью"
    ]
    assert word_mismatch(speech, word.split("Слайд 2.")[0]) == [
        "Word: разделов «Слайд N.» 1, в речи 2"
    ]
    assert words_range("720-860") == (720, 860) and words_range("800") == (720, 880)
    text = "## Слайд 1. А\nРаз два три.\n## Слайд 2. Б\nКак говорилось на лекции, четыре.\n"
    assert check(text, 2, 5, 20) == ([], [], [3, 5])
    errors, warnings, _ = check(text, 3, 1, 2, forbid="лекци")
    assert len(errors) == 2 and "длиннее цели" in warnings[0], (
        errors
    )  # число разделов, запрет; объём — мягкий
    assert check(text, 2, 20, 30)[0] == ["слов 8, нужно 20–30"] and check(
        text, 2, None, None
    )[:2] == ([], [])
    timed = "## Слайд 1. А (1 мин)\nРаз два три.\n## Слайд 2. Б (0:45)\nКак говорилось на лекции, четыре.\n"
    assert check(timed, 2, 5, 20)[1] == [
        "время в заголовках слайдов: ['## Слайд 1. А (1 мин)', '## Слайд 2. Б (0:45)']"
    ]
    assert check(timed, 2, 5, 20, strict=True)[0]
    assert check("## Слайд 1. Рост (2015–2025)\nРаз два три.\n", 1, 1, 9) == (
        [],
        [],
        [3],
    )
    print("selftest ok")


if __name__ == "__main__":
    if sys.argv[1:] == ["--selftest"]:
        selftest()
        sys.exit()
    ap = argparse.ArgumentParser()
    ap.add_argument("script")
    ap.add_argument("--slides", type=int, required=True)
    ap.add_argument(
        "--words", type=words_range, help="диапазон слов 720-860 или цель 790 (±10 %%)"
    )
    ap.add_argument("--forbid")
    ap.add_argument("--pptx")
    ap.add_argument("--docx")
    ap.add_argument(
        "--strict",
        action="store_true",
        help="учебная сдача: стоп-фразы, «—» и «ё» в Word — ошибки",
    )
    ap.add_argument("--allow", help="регулярка: разрешённый контекст стоп-фразы")
    a = ap.parse_args()
    lo, hi = a.words or (None, None)
    deck = None
    if a.pptx:
        from pptx import Presentation

        deck = Presentation(a.pptx)
    errors, warnings, per = check(
        Path(a.script).read_text(encoding="utf-8"),
        a.slides,
        lo,
        hi,
        a.forbid,
        deck,
        strict=a.strict,
        allow=a.allow,
        word=a.docx,
    )
    print(f"разделов={len(per)} слов={sum(per)} по слайдам={per}")
    for e in errors:
        print("  - ошибка:", e)
    for w in warnings:
        print("  - предупреждение:", w)
    sys.exit(1 if errors else 0)
